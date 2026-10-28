"""Extract text from PDF, DOCX, and image resumes, using OCR when needed.

Examples:
    python extract_resume.py "resume.pdf"
    python extract_resume.py "resume.png" --output extracted.txt
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import fitz  # PyMuPDF
import pytesseract
from PIL import Image, ImageOps

from config import (
    PDF_RENDER_DPI,
    PDF_TEXT_MIN_CHARS,
    SUPPORTED_IMAGE_EXTENSIONS,
    TESSERACT_CMD,
)

if TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD


def ocr_image(image: Image.Image) -> str:
    """OCR an image after light preprocessing."""
    image = ImageOps.exif_transpose(image)
    image = ImageOps.grayscale(image)
    # Improve contrast without assuming all scans need aggressive thresholding.
    image = ImageOps.autocontrast(image)
    return pytesseract.image_to_string(image, config="--psm 3").strip()


def extract_pdf(path: Path) -> tuple[str, list[str]]:
    """Extract embedded text; OCR individual pages that have little/no text."""
    sections: list[str] = []
    methods: list[str] = []

    with fitz.open(path) as document:
        for page_number, page in enumerate(document, start=1):
            embedded = page.get_text("text").strip()
            if len(embedded) >= PDF_TEXT_MIN_CHARS:
                text = embedded
                method = "embedded PDF text"
            else:
                pixmap = page.get_pixmap(
                    dpi=PDF_RENDER_DPI,
                    colorspace=fitz.csGRAY,
                    alpha=False,
                )
                image = Image.frombytes(
                    "L", [pixmap.width, pixmap.height], pixmap.samples
                )
                text = ocr_image(image)
                method = "OCR fallback" if text else "OCR fallback (no text detected)"

            sections.append(f"--- Page {page_number} ({method}) ---\n{text}")
            methods.append(f"Page {page_number}: {method}")

    return "\n\n".join(sections).strip(), methods


def extract_docx(path: Path) -> tuple[str, list[str]]:
    """Extract paragraphs and tables from a Word document."""
    from docx import Document

    document = Document(path)
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"[Table {table_index}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    result = "\n".join(parts).strip()
    return result, ["DOCX text extraction"]


def extract_image(path: Path) -> tuple[str, list[str]]:
    with Image.open(path) as image:
        text = ocr_image(image)
    return text, ["Image OCR"]


def extract_resume(file_path: str | Path) -> tuple[str, list[str]]:
    path = Path(file_path).expanduser().resolve()
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"File does not exist: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix in SUPPORTED_IMAGE_EXTENSIONS:
        return extract_image(path)

    raise ValueError(
        f"Unsupported file type '{suffix}'. Supported: PDF, DOCX, "
        f"{', '.join(sorted(SUPPORTED_IMAGE_EXTENSIONS))}"
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract text from a resume (PDF, DOCX, or image), using OCR when needed."
    )
    parser.add_argument("resume", help="Path to the resume file")
    parser.add_argument("--output", "-o", help="Optional path to save extracted text")
    args = parser.parse_args()

    try:
        text, methods = extract_resume(args.resume)
    except pytesseract.TesseractNotFoundError:
        print(
            "ERROR: Tesseract OCR was not found. Install Tesseract and either add it "
            "to PATH or set TESSERACT_CMD in config.py.",
            file=sys.stderr,
        )
        return 2
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    print("Extraction methods:")
    for method in methods:
        print(f"  - {method}")
    print("\n" + "=" * 72 + "\nEXTRACTED TEXT\n" + "=" * 72)
    print(text or "[No text extracted]")

    if args.output:
        output_path = Path(args.output).expanduser().resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(text, encoding="utf-8")
        print(f"\nSaved extracted text to: {output_path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
